"""トリミング済みデータとネイティブグラフをExcelへ出力する。"""

from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .constants import DECIMALS, JARS, METRICS, METRIC_LABELS
from .parser import to_jar_frame

HEADER_FILL = "4472C4"


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)


def _write_dataframe(ws, frame: pd.DataFrame, headers: list[str]) -> None:
    """結合元データなど、単純な表形式データを書き込む。"""
    ws.append(headers)
    _style_header(ws)
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).number_format = "yyyy/mm/dd hh:mm"


def _trim_jar(merged: pd.DataFrame, jar: str, selection: dict[str, Any]) -> pd.DataFrame:
    frame = to_jar_frame(merged, jar)
    return frame[frame["index"].between(selection["start_index"], selection["end_index"], inclusive="both")].copy()


def _write_combined_data(ws, trimmed_by_jar: dict[str, pd.DataFrame]) -> dict[str, dict[str, int]]:
    """全ジャーを、開始点を揃えた1枚のシートへ横並びで書き込む。"""
    headers = ["経過時間"]
    metric_columns: dict[str, dict[str, int]] = {}
    for jar in trimmed_by_jar:
        headers.extend([f"Jar{jar} TIME", f"Jar{jar} INDEX"] + [f"Jar{jar} {METRIC_LABELS[metric]}" for metric in METRICS])
        first_metric_col = len(headers) - len(METRICS) + 1
        metric_columns[jar] = {metric: first_metric_col + pos for pos, metric in enumerate(METRICS)}
    ws.append(headers)
    _style_header(ws)

    max_rows = max((len(frame) for frame in trimmed_by_jar.values()), default=0)
    for position in range(max_rows):
        row: list[Any] = [timedelta(minutes=5 * position)]
        for frame in trimmed_by_jar.values():
            if position < len(frame):
                item = frame.iloc[position]
                row.extend([item["time"], int(item["index"])] + [item[metric] for metric in METRICS])
            else:
                row.extend([None] * (2 + len(METRICS)))
        ws.append(row)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 14
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).number_format = "[h]:mm"
        for jar_pos, jar in enumerate(trimmed_by_jar):
            base_col = 2 + jar_pos * (2 + len(METRICS))
            ws.cell(row, base_col).number_format = "yyyy/mm/dd hh:mm"
            for metric_pos, metric in enumerate(METRICS):
                cell = ws.cell(row, base_col + 2 + metric_pos)
                cell.number_format = "0" if DECIMALS[metric] == 0 else "0." + "0" * DECIMALS[metric]
    for column in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 18
    return metric_columns


def _add_combined_scatter_charts(ws, trimmed_by_jar: dict[str, pd.DataFrame], metric_columns: dict[str, dict[str, int]]) -> None:
    """各項目について、出力対象ジャーを重ねた散布図を追加する。"""
    chart_col = get_column_letter(ws.max_column + 2)
    for metric_pos, metric in enumerate(METRICS):
        chart = ScatterChart()
        chart.title = f"{METRIC_LABELS[metric]}の推移"
        chart.y_axis.title = METRIC_LABELS[metric]
        chart.x_axis.title = "経過時間（開始点=0:00）"
        chart.x_axis.numFmt = "[h]:mm"
        chart.scatterStyle = "line"
        chart.height = 7.5
        chart.width = 14.5
        chart.style = 13
        for jar, frame in trimmed_by_jar.items():
            if frame.empty:
                continue
            max_row = len(frame) + 1
            xvalues = Reference(ws, min_col=1, min_row=2, max_row=max_row)
            yvalues = Reference(ws, min_col=metric_columns[jar][metric], min_row=2, max_row=max_row)
            series = Series(yvalues, xvalues, title=f"Jar{jar}")
            series.marker.symbol = "none"
            series.graphicalProperties.line.width = 19050
            chart.series.append(series)
        ws.add_chart(chart, f"{chart_col}{2 + metric_pos * 16}")


def build_excel(merged: pd.DataFrame, selections: dict[str, dict[str, Any]], selected_jars: list[str], parameters: dict[str, Any], include_merged: bool = False) -> bytes:
    """現在の選択区間からExcelワークブックを生成する。"""
    valid_jars = [jar for jar in selected_jars if jar in JARS]
    if not valid_jars:
        raise ValueError("出力対象ジャーを1つ以上選択してください。")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "概要"
    summary.append(["ジャー", "開始INDEX", "開始時刻", "終了INDEX", "終了時刻", "行数", "設定方法"])
    _style_header(summary)

    trimmed_by_jar: dict[str, pd.DataFrame] = {}
    for jar in valid_jars:
        trimmed = _trim_jar(merged, jar, selections[jar])
        trimmed_by_jar[jar] = trimmed
        start = trimmed.iloc[0] if not trimmed.empty else None
        end = trimmed.iloc[-1] if not trimmed.empty else None
        summary.append([jar, None if start is None else int(start["index"]), None if start is None else start["time"], None if end is None else int(end["index"]), None if end is None else end["time"], len(trimmed), selections[jar].get("mode", "manual")])

    summary.append([])
    summary.append(["自動検出パラメータ", "値"])
    for key, value in parameters.items():
        summary.append([key, value])
    summary.column_dimensions["A"].width = 24
    for column in range(2, 8):
        summary.column_dimensions[get_column_letter(column)].width = 20
    for row in range(2, 2 + len(valid_jars)):
        summary.cell(row, 3).number_format = "yyyy/mm/dd hh:mm"
        summary.cell(row, 5).number_format = "yyyy/mm/dd hh:mm"

    data_sheet = workbook.create_sheet("トリミングデータ")
    metric_columns = _write_combined_data(data_sheet, trimmed_by_jar)
    _add_combined_scatter_charts(data_sheet, trimmed_by_jar, metric_columns)

    if include_merged:
        ws = workbook.create_sheet("結合元データ")
        raw_headers = ["TIME", "INDEX"] + [f"{jar} {METRIC_LABELS[m]}" for jar in JARS for m in METRICS]
        _write_dataframe(ws, merged, raw_headers)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
