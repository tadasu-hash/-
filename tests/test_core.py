from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from core.constants import JARS, METRICS
from core.detect import detect_interval
from core.export import build_excel
from core.merge import merge_logs
from core.parser import ParsedLog, parse_logger_csv, to_jar_frame


def make_frame(indices=(100, 101, 102, 103, 104, 105, 106)):
    data = {"time": pd.date_range("2026-05-13", periods=len(indices), freq="5min"), "index": indices}
    for jar in JARS:
        for metric in METRICS:
            data[f"{jar}_{metric}"] = [0.0] * len(indices)
    return pd.DataFrame(data)


def test_parser_uses_position_and_coerces_missing():
    rows = [["[LOGGING]"] + [str(i) for i in range(1, 26)], ["TYPE"] * 26, ["表記揺れ"] * 26]
    rows.append(["2026/05/13 03:42", "35276"] + ["1.2"] * 23 + [""])
    raw = "\n".join(",".join(row) for row in rows).encode()
    parsed = parse_logger_csv(raw, "x.csv")
    assert parsed.data.loc[0, "index"] == 35276
    assert parsed.data.loc[0, "D_temp"] != parsed.data.loc[0, "D_temp"]  # NaN


def test_merge_later_file_wins_and_keeps_gap():
    first, second = make_frame((1, 3)), make_frame((3, 4))
    first.loc[first["index"] == 3, "A_pump"] = 1
    second.loc[second["index"] == 3, "A_pump"] = 9
    merged, summary = merge_logs([ParsedLog("a.csv", first, [], []), ParsedLog("b.csv", second, [], [])])
    assert merged["index"].tolist() == [1, 3, 4]
    assert merged.loc[merged["index"] == 3, "A_pump"].item() == 9
    assert summary.duplicates_resolved == 1
    assert summary.missing_indices == 1


def test_detection_ignores_single_spike_and_falls_back_end():
    frame = to_jar_frame(make_frame(), "A")
    frame["pump"] = [0, 100, 0, 0, 17.9, 17.9, 17.9]
    frame["temp"] = [25] * 7
    result = detect_interval(frame)
    assert result.start_pos == 3
    assert result.end_pos == 6
    assert result.start_detected and not result.end_detected


def test_detection_stuck_sensor_falls_back_all():
    frame = to_jar_frame(make_frame(), "B")
    frame["pump"] = 0
    frame["temp"] = -3.2
    result = detect_interval(frame)
    assert (result.start_pos, result.end_pos) == (0, len(frame) - 1)
    assert result.warning


def test_excel_only_selected_jars_and_native_charts():
    merged = make_frame()
    selections = {jar: {"start_index": 101, "end_index": 105, "mode": "manual"} for jar in JARS}
    output = build_excel(merged, selections, ["A", "C"], {"START_THRESHOLD": 5})
    book = load_workbook(BytesIO(output))
    assert book.sheetnames == ["概要", "JarA", "JarC"]
    assert len(book["JarA"]._charts) == 6
    assert book["JarA"].max_row == 6
